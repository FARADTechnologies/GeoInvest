"""Parse the mass-valuation Excel template into listing rows.

The team's template (`nümunə.xlsx`) has one listing per row with these AZ
headers (order-independent — matched by name):

    Kateqoriya | Çıxarış | Yaşayış kompleksi | Kompleks adı | Təmir statusu |
    Sahə | Otaq sayı | Yerləşdiyi mərtəbə | Binanın mərtəbə sayı | Ünvan

There are no coordinates in the file, only an address; the frontend geocodes
each address before calling the predict model (same as the single form).
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

# Normalised header (strip + lower) → our field name.
_HEADER_MAP = {
    "kateqoriya": "type",
    "çıxarış": "extract",
    "yaşayış kompleksi": "is_residence",
    "kompleks adı": "residence",
    "təmir statusu": "repair",
    "təmir": "repair",
    "sahə": "area",
    "otaq sayı": "rooms",
    "yerləşdiyi mərtəbə": "floor",
    "binanın mərtəbə sayı": "total_floors",
    "ünvan": "address",
}

_TRUE_WORDS = {"var", "bəli", "beli", "bali", "yes", "true", "1", "hə", "he"}


def _num(v) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(str(v).replace(",", ".").strip())
    except (TypeError, ValueError):
        return None


def _int(v) -> int | None:
    n = _num(v)
    return int(round(n)) if n is not None else None


def _str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def parse_listings(data: bytes) -> list[dict]:
    """Return a list of listing dicts from the uploaded .xlsx bytes.

    read_only + data_only means no formulas/macros are evaluated — the file is
    only read as data, never executed.
    """
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [(_str(c) or "").lower() for c in rows[0]]
    idx: dict[str, int] = {}
    for i, h in enumerate(header):
        field = _HEADER_MAP.get(h)
        if field and field not in idx:
            idx[field] = i

    def cell(row, field):
        i = idx.get(field)
        return row[i] if i is not None and i < len(row) else None

    out: list[dict] = []
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        res_raw = (_str(cell(row, "is_residence")) or "").lower()
        out.append(
            {
                "type": _str(cell(row, "type")),
                "extract": _str(cell(row, "extract")),
                "is_residence": res_raw in _TRUE_WORDS,
                "residence": _str(cell(row, "residence")),
                "repair": _str(cell(row, "repair")),
                "area": _num(cell(row, "area")),
                "rooms": _int(cell(row, "rooms")),
                "floor": _int(cell(row, "floor")),
                "total_floors": _int(cell(row, "total_floors")),
                "address": _str(cell(row, "address")),
            }
        )
    return out
